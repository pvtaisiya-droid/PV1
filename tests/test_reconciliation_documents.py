from datetime import date
from io import BytesIO
from types import SimpleNamespace

from docx import Document

from app import reconciliation_documents
from app.reconciliation import render_reconciliation_email


def reconciliation_fixture(language: str = "ru"):
    substance = SimpleNamespace(inn_name="Acetylsalicylic acid", substance_name="Ацетилсалициловая кислота")
    substance_link = SimpleNamespace(
        substance=substance,
        is_deleted=False,
        is_active=True,
    )
    certificate = SimpleNamespace(
        registration_number="ЛП-001234",
        country_code="RU",
        dosage_form="таблетки",
        strength="100 мг",
        status="active",
        is_deleted=False,
    )
    product = SimpleNamespace(
        product_name="Тестовый препарат",
        authorization_number=None,
        authorization_country_code=None,
        dosage_form="таблетки",
        strength="100 мг",
        registration_certificates=[certificate],
        substance_links=[substance_link],
        is_deleted=False,
        dosage_form_name=lambda requested_language: "tablets" if requested_language == "en" else "таблетки",
    )
    partner = SimpleNamespace(
        partner_name="Тестовый партнёр",
        partner_code="TP-001",
        products=[product],
    )
    our_item = SimpleNamespace(
        source_side="our_company",
        internal_case_number="2026-001",
        partner_case_number=None,
        partner_case_id=None,
        receipt_date_our_company=date(2026, 8, 3),
        transfer_date_our_company=date(2026, 8, 4),
        receipt_date_partner=None,
        transfer_date_partner=None,
        adverse_event="Описание сообщения",
        short_description=None,
        reconciliation_status="matched",
    )
    partner_item = SimpleNamespace(
        source_side="partner",
        internal_case_number=None,
        partner_case_number="PARTNER-001",
        partner_case_id=None,
        receipt_date_our_company=None,
        transfer_date_our_company=None,
        receipt_date_partner=date(2026, 8, 7),
        transfer_date_partner=date(2026, 8, 8),
        adverse_event="Partner message description",
        short_description=None,
        reconciliation_status="matched",
    )
    reconciliation = SimpleNamespace(
        id="test-reconciliation",
        language=language,
        partner=partner,
        period_start=date(2026, 8, 1),
        period_end=date(2026, 8, 31),
        items=[our_item, partner_item],
        contact_name="Contact Person",
        contact_email="contact@example.com",
        prepared_by="PV Specialist",
        reconciliation_status="draft",
    )
    return reconciliation, [product]


def docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def test_reconciliation_docx_uses_approved_russian_form():
    reconciliation, products = reconciliation_fixture("ru")
    content = reconciliation_documents.build_reconciliation_document_bytes(
        reconciliation,
        document_format="docx",
        products=products,
    )
    text = docx_text(content)

    assert "ФОРМА СВЕРКИ СООБЩЕНИЙ ПО БЕЗОПАСНОСТИ" in text
    assert "№ РУ" in text
    assert "Торговое наименование" in text
    assert "МНН" in text
    assert "Лекарственная форма и дозировка" in text
    assert "Номер сообщения" in text
    assert "Описание" in text
    assert "Дата получения" in text
    assert "Дата передачи другой стороне" in text
    assert "ООО «АРС»" in text
    assert "Дата" in text
    assert "Подпись" in text
    assert "СОП-" not in text
    assert "Дата введения" not in text
    assert "Срок действия" not in text


def test_reconciliation_docx_supports_english_version():
    reconciliation, products = reconciliation_fixture("en")
    content = reconciliation_documents.build_reconciliation_document_bytes(
        reconciliation,
        document_format="docx",
        products=products,
    )
    text = docx_text(content)

    assert "SAFETY MESSAGES RECONCILIATION FORM" in text
    assert "under the pharmacovigilance agreement" not in text
    assert "Marketing authorisation No." in text
    assert "Trade name" in text
    assert "Date of receipt Customer" in text
    assert "Date of receipt ARS LLC" in text
    assert "Date transmitted to the other party" not in text
    assert "ARS, LLC" in text
    assert "Signature" in text


def test_reconciliation_document_uses_only_current_registration_certificate():
    reconciliation, products = reconciliation_fixture("en")
    product = products[0]
    product.registration_certificates = [
        SimpleNamespace(
            registration_number="OLD-REPLACED",
            country_code="RU",
            dosage_form="tablets",
            strength="50 mg",
            status="replaced",
            registration_date=date(2024, 1, 1),
            valid_until=None,
            is_deleted=False,
        ),
        SimpleNamespace(
            registration_number="OLD-EXPIRED",
            country_code="RU",
            dosage_form="tablets",
            strength="75 mg",
            status="active",
            registration_date=date(2025, 1, 1),
            valid_until=date(2025, 12, 31),
            is_deleted=False,
        ),
        SimpleNamespace(
            registration_number="CURRENT-RU",
            country_code="RU",
            dosage_form="tablets",
            strength="100 mg",
            status="registration_active",
            registration_date=date(2026, 1, 1),
            valid_until=None,
            is_deleted=False,
        ),
    ]

    content = reconciliation_documents.build_reconciliation_document_bytes(
        reconciliation,
        document_format="docx",
        products=products,
    )
    text = docx_text(content)

    assert "CURRENT-RU" in text
    assert "OLD-REPLACED" not in text
    assert "OLD-EXPIRED" not in text


def test_pdf_generation_creates_docx_and_pdf_pair(tmp_path, monkeypatch):
    reconciliation, products = reconciliation_fixture("ru")
    monkeypatch.setattr(reconciliation_documents, "RECONCILIATION_DOCUMENT_DIR", tmp_path)

    document_info = reconciliation_documents.generate_reconciliation_document_file(
        reconciliation,
        document_format="pdf",
        products=products,
    )

    pdf_path = tmp_path / document_info["pdf_filename"]
    docx_path = tmp_path / document_info["docx_filename"]
    assert document_info["document_format"] == "pdf"
    assert document_info["mime_type"] == "application/pdf"
    assert pdf_path.read_bytes().startswith(b"%PDF")
    assert docx_path.read_bytes().startswith(b"PK")
    assert pdf_path.stat().st_size > 10_000
    assert docx_path.stat().st_size > 10_000


def test_reconciliation_email_uses_selected_language():
    reconciliation, _ = reconciliation_fixture("ru")
    reconciliation.contact = SimpleNamespace(first_name="Анна")
    reconciliation.contact_name = "Иванова Анна"
    reconciliation.document_filename = "reconciliation.pdf"
    reconciliation.email_subject = None
    reconciliation.email_body = None

    russian_email = render_reconciliation_email(reconciliation)

    assert "Уважаемая Анна," in russian_email["body"]
    assert "с 1 августа 2026 года по 31 августа 2026 года" in russian_email["body"]
    assert "Таисия Воркунова" in russian_email["body"]
    assert "УЛФ" in russian_email["body"]
    assert "Dear " not in russian_email["body"]
    assert "Safety Messages Reconciliation Form" not in russian_email["subject"]

    reconciliation.language = "en"
    english_email = render_reconciliation_email(reconciliation)

    assert "Dear Анна," in english_email["body"]
    assert "from 1 August 2026 to 31 August 2026" in english_email["body"]
    assert "Taisiya Vorkunova" in english_email["body"]
    assert "QPPV" in english_email["body"]
    assert "Уважаемая " not in english_email["body"]
    assert "Safety Messages Reconciliation Form" in english_email["subject"]
    assert "tvorkunova@pharmrussia.com" in english_email["body"]
