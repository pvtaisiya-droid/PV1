from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.reconciliation_excel import build_reconciliation_workbook


from app.config import get_settings

RECONCILIATION_DOCUMENT_DIR = get_settings().storage_root / "reconciliations"
BRAND_LOGO_PATH = Path(__file__).resolve().parent / "static" / "brand" / "ars-pharmrussia-logo.png"
SUPPORTED_RECONCILIATION_FORMATS = {"xlsx", "docx", "pdf"}

INK = "1F2933"
MUTED = "5E6B75"
ACCENT = "1F4D78"
TABLE_FILL = "F2F4F7"
TABLE_BORDER = "AAB4BD"
FORM_WIDTH_DXA = 10200
TABLE_INDENT_DXA = 120
CELL_MARGIN_VERTICAL_DXA = 80
CELL_MARGIN_HORIZONTAL_DXA = 120
ACTIVE_REGISTRATION_STATUSES = {"active", "registration_active", "authorized"}


TEXT = {
    "ru": {
        "title": "ФОРМА СВЕРКИ СООБЩЕНИЙ ПО БЕЗОПАСНОСТИ",
        "subtitle": "в рамках соглашения по фармаконадзору",
        "partner": "Партнёр",
        "partner_code": "Код партнёра",
        "period": "Период сверки",
        "products": "Лекарственные препараты",
        "registration": "№ РУ",
        "trade_name": "Торговое наименование",
        "inn": "МНН",
        "form_strength": "Лекарственная форма и дозировка",
        "no_products": "Лекарственные препараты отсутствуют",
        "from_ars": "1. Сообщения по безопасности, полученные ООО «АРС» и переданные партнёру",
        "from_partner": "2. Сообщения по безопасности, полученные партнёром и переданные ООО «АРС»",
        "message_number": "Номер сообщения",
        "description": "Описание",
        "received_date": "Дата получения",
        "transferred_date": "Дата передачи другой стороне",
        "none": "Сообщения отсутствуют",
        "date": "Дата",
        "signature": "Подпись",
        "ars": "ООО «АРС»",
        "page": "Страница",
        "of": "из",
    },
    "en": {
        "title": "SAFETY MESSAGES RECONCILIATION FORM",
        "subtitle": "",
        "partner": "Partner",
        "partner_code": "Partner code",
        "period": "Reconciliation period",
        "products": "Medicinal products",
        "registration": "Marketing authorisation No.",
        "trade_name": "Trade name",
        "inn": "INN",
        "form_strength": "Pharmaceutical form and strength",
        "no_products": "No medicinal products",
        "from_ars": "1. Safety messages received by ARS, LLC and transmitted to the Partner",
        "from_partner": "2. Safety messages received by the Partner and transmitted to ARS, LLC",
        "message_number": "Message number",
        "description": "Description",
        "received_by_ars": "Date of receipt ARS LLC",
        "received_by_customer": "Date of receipt Customer",
        "none": "No safety messages",
        "date": "Date",
        "signature": "Signature",
        "ars": "ARS, LLC",
        "page": "Page",
        "of": "of",
    },
}


def generate_reconciliation_document_file(
    reconciliation,
    *,
    document_format: str = "pdf",
    products: list[Any] | None = None,
    created_paths: list[Path] | None = None,
    unique_name: bool = False,
) -> dict[str, Any]:
    document_format = normalize_document_format(document_format)
    RECONCILIATION_DOCUMENT_DIR.mkdir(parents=True, exist_ok=True)
    language = normalize_language(getattr(reconciliation, "language", None))
    stem = f"partner_reconciliation_{reconciliation.id}_{language}"
    if unique_name:
        stem += f"_{uuid4().hex}"

    if document_format == "xlsx":
        filename = f"{stem}.xlsx"
        path = RECONCILIATION_DOCUMENT_DIR / filename
        content = build_reconciliation_document_bytes(
            reconciliation,
            document_format="xlsx",
            products=products,
        )
        if created_paths is not None:
            created_paths.append(path)
        path.write_bytes(content)
        return document_info(path, "xlsx", len(content))

    docx_bytes = build_reconciliation_document_bytes(
        reconciliation,
        document_format="docx",
        products=products,
    )
    pdf_bytes = build_reconciliation_document_bytes(
        reconciliation,
        document_format="pdf",
        products=products,
    )
    docx_path = RECONCILIATION_DOCUMENT_DIR / f"{stem}.docx"
    pdf_path = RECONCILIATION_DOCUMENT_DIR / f"{stem}.pdf"
    if created_paths is not None:
        created_paths.append(docx_path)
    docx_path.write_bytes(docx_bytes)
    if created_paths is not None:
        created_paths.append(pdf_path)
    pdf_path.write_bytes(pdf_bytes)

    primary_path = pdf_path if document_format == "pdf" else docx_path
    primary_bytes = pdf_bytes if document_format == "pdf" else docx_bytes
    info = document_info(primary_path, document_format, len(primary_bytes))
    info.update(
        {
            "docx_path": str(docx_path),
            "docx_filename": docx_path.name,
            "pdf_path": str(pdf_path),
            "pdf_filename": pdf_path.name,
        }
    )
    return info


def document_info(path: Path, document_format: str, size: int) -> dict[str, Any]:
    return {
        "document_path": str(path),
        "document_filename": path.name,
        "document_format": document_format,
        "file_size_bytes": size,
        "mime_type": mime_type_for_document_format(document_format),
    }


def build_reconciliation_document_bytes(
    reconciliation,
    *,
    document_format: str = "pdf",
    products: list[Any] | None = None,
) -> bytes:
    document_format = normalize_document_format(document_format)
    our_items = [item for item in reconciliation.items if item.source_side == "our_company"]
    partner_items = [item for item in reconciliation.items if item.source_side == "partner"]
    if document_format == "xlsx":
        discrepancy_items = [
            item
            for item in reconciliation.items
            if item.reconciliation_status not in {"matched", "confirmed"}
        ]
        return build_reconciliation_workbook(
            partner_name=reconciliation.partner.partner_name if reconciliation.partner else "",
            period_start=reconciliation.period_start,
            period_end=reconciliation.period_end,
            generated_at=date.today(),
            contact_name=reconciliation.contact_name,
            contact_email=reconciliation.contact_email,
            prepared_by=reconciliation.prepared_by,
            status=reconciliation.reconciliation_status,
            our_items=our_items,
            partner_items=partner_items,
            discrepancy_items=discrepancy_items,
            language=reconciliation.language,
        )

    resolved_products = products if products is not None else products_from_reconciliation(reconciliation)
    if document_format == "docx":
        return build_reconciliation_docx(
            reconciliation,
            products=resolved_products,
            our_items=our_items,
            partner_items=partner_items,
        )
    return build_reconciliation_pdf(
        reconciliation,
        products=resolved_products,
        our_items=our_items,
        partner_items=partner_items,
    )


def normalize_document_format(document_format: str | None) -> str:
    value = (document_format or "pdf").strip().lower()
    if value not in SUPPORTED_RECONCILIATION_FORMATS:
        raise ValueError("Unsupported reconciliation document format.")
    return value


def normalize_language(language: str | None) -> str:
    return "en" if (language or "").lower() == "en" else "ru"


def mime_type_for_document_format(document_format: str) -> str:
    if document_format == "docx":
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if document_format == "pdf":
        return "application/pdf"
    return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_reconciliation_docx(
    reconciliation,
    *,
    products: list[Any],
    our_items: list[Any],
    partner_items: list[Any],
) -> bytes:
    language = normalize_language(reconciliation.language)
    labels = TEXT[language]
    document = Document()
    configure_document(document, labels)
    add_docx_header(document, language)
    add_docx_title(document, labels)
    add_docx_metadata(document, reconciliation, labels)
    add_docx_section_heading(document, labels["products"])
    add_docx_product_table(document, products, language, labels)
    add_docx_section_heading(document, labels["from_ars"])
    add_docx_message_table(document, our_items, "our_company", language, labels)
    add_docx_section_heading(document, labels["from_partner"])
    add_docx_message_table(document, partner_items, "partner", language, labels)
    add_docx_signature_block(document, labels["ars"], labels)
    partner_name = reconciliation.partner.partner_name if reconciliation.partner else labels["partner"]
    add_docx_signature_block(document, partner_name, labels)

    document.core_properties.title = labels["title"]
    document.core_properties.subject = partner_name
    document.core_properties.author = labels["ars"]
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def configure_document(document: Document, labels: dict[str, str]) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading = document.styles["Heading 2"]
    heading.font.name = "Arial"
    heading._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    heading.font.size = Pt(11.5)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(ACCENT)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(5)
    heading.paragraph_format.keep_with_next = True

    add_docx_footer(section, labels)


def add_docx_header(document: Document, language: str) -> None:
    section = document.sections[0]
    header = section.header
    table = header.add_table(rows=1, cols=2, width=Cm(18))
    set_table_geometry(table, [1600, 8600], indent_dxa=0)
    remove_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    if BRAND_LOGO_PATH.exists():
        paragraph = table.cell(0, 0).paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run()
        run.add_picture(str(BRAND_LOGO_PATH), width=Cm(1.25))
    company = TEXT[language]["ars"]
    set_cell_text(table.cell(0, 1), company, bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)


def add_docx_footer(section, labels: dict[str, str]) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(f"{labels['page']} ")
    set_run_font(run, size=8, color=MUTED)
    add_field(paragraph, "PAGE")
    run = paragraph.add_run(f" {labels['of']} ")
    set_run_font(run, size=8, color=MUTED)
    add_field(paragraph, "NUMPAGES")


def add_field(paragraph, field_name: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {field_name} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)


def add_docx_title(document: Document, labels: dict[str, str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(labels["title"])
    set_run_font(run, size=15, bold=True, color=ACCENT)

    if labels["subtitle"]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(12)
        run = paragraph.add_run(labels["subtitle"])
        set_run_font(run, size=10, color=MUTED)
    else:
        paragraph.paragraph_format.space_after = Pt(12)


def add_docx_metadata(document: Document, reconciliation, labels: dict[str, str]) -> None:
    partner = reconciliation.partner
    rows = [
        (labels["partner"], partner.partner_name if partner else ""),
        (labels["partner_code"], partner.partner_code if partner else ""),
        (
            labels["period"],
            f"{format_date(reconciliation.period_start)} - {format_date(reconciliation.period_end)}",
        ),
    ]
    table = document.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [2200, 8000])
    remove_table_borders(table)
    for index, (label, text) in enumerate(rows):
        set_cell_text(table.cell(index, 0), label, bold=True, size=9.5, color=MUTED)
        set_cell_text(table.cell(index, 1), text, bold=True, size=10.5)
    add_spacing_after_table(document, 2)


def add_docx_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Heading 2")
    paragraph.add_run(text)


def add_docx_product_table(
    document: Document,
    products: list[Any],
    language: str,
    labels: dict[str, str],
) -> None:
    headers = [labels["registration"], labels["trade_name"], labels["inn"], labels["form_strength"]]
    rows = [product_row(product, language) for product in products]
    if not rows:
        rows = [["-", labels["no_products"], "-", "-"]]
    add_docx_table(document, headers, rows, [1800, 2400, 2700, 3300], center_columns={0})


def add_docx_message_table(
    document: Document,
    items: list[Any],
    source_side: str,
    language: str,
    labels: dict[str, str],
) -> None:
    headers = message_table_headers(labels, source_side, language)
    rows = [message_row(item, source_side) for item in items]
    if not rows:
        rows = [["-", labels["none"], "-", "-"]]
    add_docx_table(document, headers, rows, [1700, 4700, 1900, 1900], center_columns={0, 2, 3})


def message_table_headers(labels: dict[str, str], source_side: str, language: str) -> list[str]:
    if language == "en":
        receipt_headers = (
            [labels["received_by_ars"], labels["received_by_customer"]]
            if source_side == "our_company"
            else [labels["received_by_customer"], labels["received_by_ars"]]
        )
    else:
        receipt_headers = [labels["received_date"], labels["transferred_date"]]
    return [labels["message_number"], labels["description"], *receipt_headers]


def add_docx_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
    *,
    center_columns: set[int],
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths_dxa)
    set_repeat_table_header(table.rows[0])
    for column, header in enumerate(headers):
        set_cell_text(
            table.rows[0].cells[column],
            header,
            bold=True,
            size=9,
            color=INK,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            fill=TABLE_FILL,
        )
    for row_values in rows:
        row = table.add_row()
        for column, cell_value in enumerate(row_values):
            alignment = WD_ALIGN_PARAGRAPH.CENTER if column in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[column], cell_value, size=9, align=alignment)
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    add_spacing_after_table(document, 4)


def add_docx_signature_block(document: Document, party_name: str, labels: dict[str, str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(party_name)
    set_run_font(run, size=10.5, bold=True)

    table = document.add_table(rows=2, cols=2)
    set_table_geometry(table, [5100, 5100])
    remove_table_borders(table)
    for column in range(2):
        line = table.cell(0, column).paragraphs[0]
        line.paragraph_format.space_before = Pt(8)
        line.paragraph_format.space_after = Pt(0)
        line.add_run(" ")
        set_paragraph_bottom_border(line)
    set_cell_text(table.cell(1, 0), labels["date"], size=8, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(1, 1), labels["signature"], size=8, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)


def set_table_geometry(table, widths_dxa: list[int], *, indent_dxa: int = TABLE_INDENT_DXA) -> None:
    table.autofit = False
    total = sum(widths_dxa)
    table_pr = table._tbl.tblPr
    set_or_replace_measure(table_pr, "w:tblW", total, "dxa")
    set_or_replace_measure(table_pr, "w:tblInd", indent_dxa, "dxa")
    layout = table_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_pr = cell._tc.get_or_add_tcPr()
            set_or_replace_measure(cell_pr, "w:tcW", width, "dxa")
            set_cell_margins(cell)


def set_or_replace_measure(parent, tag: str, width: int, measure_type: str) -> None:
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        parent.append(element)
    element.set(qn("w:w"), str(width))
    element.set(qn("w:type"), measure_type)


def set_cell_margins(cell) -> None:
    cell_pr = cell._tc.get_or_add_tcPr()
    margins = cell_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_pr.append(margins)
    for side, width in (
        ("top", CELL_MARGIN_VERTICAL_DXA),
        ("bottom", CELL_MARGIN_VERTICAL_DXA),
        ("start", CELL_MARGIN_HORIZONTAL_DXA),
        ("end", CELL_MARGIN_HORIZONTAL_DXA),
    ):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(width))
        node.set(qn("w:type"), "dxa")


def set_cell_text(
    cell,
    text: object,
    *,
    bold: bool = False,
    size: float = 9,
    color: str = INK,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    fill: str | None = None,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(str(text or ""))
    set_run_font(run, size=size, bold=bold, color=color)
    if fill:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shading)


def set_run_font(run, *, size: float, bold: bool = False, color: str = INK) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row) -> None:
    row_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row_pr.append(header)


def set_table_borders(table) -> None:
    table_pr = table._tbl.tblPr
    borders = table_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), TABLE_BORDER)


def remove_table_borders(table) -> None:
    table_pr = table._tbl.tblPr
    borders = table_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_paragraph_bottom_border(paragraph) -> None:
    paragraph_pr = paragraph._p.get_or_add_pPr()
    borders = paragraph_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        paragraph_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), INK)
    bottom.set(qn("w:space"), "1")
    borders.append(bottom)


def add_spacing_after_table(document: Document, points: float) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(points)


def build_reconciliation_pdf(
    reconciliation,
    *,
    products: list[Any],
    our_items: list[Any],
    partner_items: list[Any],
) -> bytes:
    language = normalize_language(reconciliation.language)
    labels = TEXT[language]
    regular_font, bold_font = register_pdf_fonts()
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=17 * mm,
        bottomMargin=16 * mm,
        title=labels["title"],
        author=labels["ars"],
    )
    styles = pdf_styles(regular_font, bold_font)
    story: list[Any] = []

    header_content: list[Any] = []
    if BRAND_LOGO_PATH.exists():
        header_content.append(Image(str(BRAND_LOGO_PATH), width=14 * mm, height=14 * mm))
    else:
        header_content.append(Paragraph("ARS", styles["meta_bold"]))
    header = Table(
        [[header_content[0], Paragraph(escape(labels["ars"]), styles["meta_right"]) ]],
        colWidths=[22 * mm, 158 * mm],
    )
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story.extend([header, Spacer(1, 4 * mm)])
    story.append(Paragraph(escape(labels["title"]), styles["title"]))
    if labels["subtitle"]:
        story.append(Paragraph(escape(labels["subtitle"]), styles["subtitle"]))
    else:
        story.append(Spacer(1, 3 * mm))

    partner = reconciliation.partner
    metadata = [
        [labels["partner"], partner.partner_name if partner else ""],
        [labels["partner_code"], partner.partner_code if partner else ""],
        [labels["period"], f"{format_date(reconciliation.period_start)} - {format_date(reconciliation.period_end)}"],
    ]
    story.append(pdf_metadata_table(metadata, styles))
    story.append(Spacer(1, 3 * mm))

    story.append(Paragraph(escape(labels["products"]), styles["heading"]))
    product_data = [[labels["registration"], labels["trade_name"], labels["inn"], labels["form_strength"]]]
    product_rows = [product_row(product, language) for product in products]
    product_data.extend(product_rows or [["-", labels["no_products"], "-", "-"]])
    story.append(pdf_data_table(product_data, [32 * mm, 42 * mm, 48 * mm, 58 * mm], styles, {0}))

    for heading, items, side in (
        (labels["from_ars"], our_items, "our_company"),
        (labels["from_partner"], partner_items, "partner"),
    ):
        story.append(Paragraph(escape(heading), styles["heading"]))
        rows = [message_row(item, side) for item in items]
        message_data = [message_table_headers(labels, side, language)]
        message_data.extend(rows or [["-", labels["none"], "-", "-"]])
        story.append(pdf_data_table(message_data, [30 * mm, 84 * mm, 33 * mm, 33 * mm], styles, {0, 2, 3}))

    story.append(Spacer(1, 5 * mm))
    story.append(pdf_signature_block(labels["ars"], labels, styles))
    story.append(Spacer(1, 5 * mm))
    partner_name = partner.partner_name if partner else labels["partner"]
    story.append(pdf_signature_block(partner_name, labels, styles))

    document.build(story, onFirstPage=lambda canvas, doc: draw_pdf_footer(canvas, doc, labels, regular_font), onLaterPages=lambda canvas, doc: draw_pdf_footer(canvas, doc, labels, regular_font))
    return output.getvalue()


def register_pdf_fonts() -> tuple[str, str]:
    regular_name = "ARSArial"
    bold_name = "ARSArialBold"
    if regular_name in pdfmetrics.getRegisteredFontNames():
        return regular_name, bold_name
    font_dir = Path("C:/Windows/Fonts")
    regular_path = font_dir / "arial.ttf"
    bold_path = font_dir / "arialbd.ttf"
    if regular_path.exists() and bold_path.exists():
        pdfmetrics.registerFont(TTFont(regular_name, str(regular_path)))
        pdfmetrics.registerFont(TTFont(bold_name, str(bold_path)))
        return regular_name, bold_name
    return "Helvetica", "Helvetica-Bold"


def pdf_styles(regular_font: str, bold_font: str) -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "body": ParagraphStyle(
            "ARSBody",
            parent=base["Normal"],
            fontName=regular_font,
            fontSize=8.6,
            leading=10.2,
            textColor=colors.HexColor(f"#{INK}"),
            spaceAfter=0,
        ),
        "body_center": ParagraphStyle(
            "ARSBodyCenter",
            parent=base["Normal"],
            fontName=regular_font,
            fontSize=8.6,
            leading=10.2,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{INK}"),
        ),
        "table_header": ParagraphStyle(
            "ARSTableHeader",
            parent=base["Normal"],
            fontName=bold_font,
            fontSize=8.2,
            leading=9.4,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{INK}"),
        ),
        "title": ParagraphStyle(
            "ARSTitle",
            parent=base["Normal"],
            fontName=bold_font,
            fontSize=14,
            leading=17,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{ACCENT}"),
            spaceAfter=2 * mm,
        ),
        "subtitle": ParagraphStyle(
            "ARSSubtitle",
            parent=base["Normal"],
            fontName=regular_font,
            fontSize=9,
            leading=11,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=5 * mm,
        ),
        "heading": ParagraphStyle(
            "ARSHeading",
            parent=base["Normal"],
            fontName=bold_font,
            fontSize=10,
            leading=12,
            textColor=colors.HexColor(f"#{ACCENT}"),
            spaceBefore=4 * mm,
            spaceAfter=2 * mm,
            keepWithNext=True,
        ),
        "meta": ParagraphStyle(
            "ARSMeta",
            parent=base["Normal"],
            fontName=regular_font,
            fontSize=9,
            leading=11,
            textColor=colors.HexColor(f"#{INK}"),
        ),
        "meta_bold": ParagraphStyle(
            "ARSMetaBold",
            parent=base["Normal"],
            fontName=bold_font,
            fontSize=9,
            leading=11,
            textColor=colors.HexColor(f"#{INK}"),
        ),
        "meta_right": ParagraphStyle(
            "ARSMetaRight",
            parent=base["Normal"],
            fontName=bold_font,
            fontSize=9,
            leading=11,
            alignment=2,
            textColor=colors.HexColor(f"#{INK}"),
        ),
        "signature": ParagraphStyle(
            "ARSSignature",
            parent=base["Normal"],
            fontName=bold_font,
            fontSize=9.5,
            leading=11,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{INK}"),
        ),
        "signature_label": ParagraphStyle(
            "ARSSignatureLabel",
            parent=base["Normal"],
            fontName=regular_font,
            fontSize=7.5,
            leading=9,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{MUTED}"),
        ),
    }


def pdf_metadata_table(rows: list[list[str]], styles: dict[str, ParagraphStyle]) -> Table:
    data = [
        [Paragraph(escape(label), styles["meta_bold"]), Paragraph(escape(value), styles["meta"])]
        for label, value in rows
    ]
    table = Table(data, colWidths=[38 * mm, 142 * mm], hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2 * mm),
                ("TOPPADDING", (0, 0), (-1, -1), 1.1 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1.1 * mm),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(f"#{TABLE_FILL}")),
            ]
        )
    )
    return table


def pdf_data_table(
    rows: list[list[str]],
    widths: list[float],
    styles: dict[str, ParagraphStyle],
    center_columns: set[int],
) -> Table:
    data: list[list[Paragraph]] = []
    for row_index, row in enumerate(rows):
        rendered_row = []
        for column, text in enumerate(row):
            style = styles["table_header"] if row_index == 0 else (styles["body_center"] if column in center_columns else styles["body"])
            rendered_row.append(Paragraph(escape(str(text or "")), style))
        data.append(rendered_row)
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor(f"#{TABLE_BORDER}")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{TABLE_FILL}")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2 * mm),
                ("TOPPADDING", (0, 0), (-1, -1), 1.5 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5 * mm),
            ]
        )
    )
    return table


def pdf_signature_block(
    party_name: str,
    labels: dict[str, str],
    styles: dict[str, ParagraphStyle],
) -> KeepTogether:
    table = Table(
        [
            [Paragraph(escape(party_name), styles["signature"]), ""],
            ["____________________________", "____________________________"],
            [Paragraph(escape(labels["date"]), styles["signature_label"]), Paragraph(escape(labels["signature"]), styles["signature_label"])],
        ],
        colWidths=[90 * mm, 90 * mm],
        hAlign="LEFT",
    )
    table.setStyle(
        TableStyle(
            [
                ("SPAN", (0, 0), (1, 0)),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
                ("TOPPADDING", (0, 0), (-1, -1), 1 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0.5 * mm),
            ]
        )
    )
    return KeepTogether([table])


def draw_pdf_footer(canvas, document, labels: dict[str, str], font_name: str) -> None:
    canvas.saveState()
    canvas.setFont(font_name, 7.5)
    canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
    page_text = f"{labels['page']} {document.page}"
    canvas.drawRightString(A4[0] - 15 * mm, 8 * mm, page_text)
    canvas.restoreState()


def products_from_reconciliation(reconciliation) -> list[Any]:
    partner = getattr(reconciliation, "partner", None)
    products = list(getattr(partner, "products", []) or []) if partner else []
    return sorted(
        [
            product
            for product in products
            if not getattr(product, "is_deleted", False)
            and getattr(product, "is_active", True)
        ],
        key=lambda product: product.product_name or "",
    )


def product_row(product: Any, language: str) -> list[str]:
    certificate = preferred_registration_certificate(product)
    registration_number = (
        getattr(certificate, "registration_number", None)
        if certificate
        else getattr(product, "authorization_number", None)
    ) or "-"
    country_code = (
        getattr(certificate, "country_code", None)
        if certificate
        else getattr(product, "authorization_country_code", None)
    )
    if country_code:
        registration_number = f"{registration_number} ({country_code})"

    inn_names: list[str] = []
    for link in getattr(product, "substance_links", []) or []:
        if getattr(link, "is_deleted", False) or not getattr(link, "is_active", True):
            continue
        substance = getattr(link, "substance", None)
        if not substance:
            continue
        name = getattr(substance, "inn_name", None) or getattr(substance, "substance_name", None)
        if name and name not in inn_names:
            inn_names.append(name)

    dosage_form = getattr(certificate, "dosage_form", None) if certificate else None
    if not dosage_form:
        dosage_method = getattr(product, "dosage_form_name", None)
        dosage_form = dosage_method(language) if callable(dosage_method) else getattr(product, "dosage_form", None)
    strength = (getattr(certificate, "strength", None) if certificate else None) or getattr(product, "strength", None)
    form_strength = ", ".join(part for part in (dosage_form, strength) if part) or "-"
    return [
        registration_number,
        getattr(product, "product_name", None) or "-",
        " + ".join(inn_names) or "-",
        form_strength,
    ]


def preferred_registration_certificate(product: Any):
    today = date.today()
    certificates = [
        certificate
        for certificate in (getattr(product, "registration_certificates", []) or [])
        if not getattr(certificate, "is_deleted", False)
        and (getattr(certificate, "status", "active") or "active").lower() in ACTIVE_REGISTRATION_STATUSES
        and (
            not getattr(certificate, "registration_date", None)
            or getattr(certificate, "registration_date") <= today
        )
        and (
            not getattr(certificate, "valid_until", None)
            or getattr(certificate, "valid_until") >= today
        )
    ]
    return max(
        certificates,
        key=lambda certificate: getattr(certificate, "registration_date", None) or date.min,
        default=None,
    )


def message_row(item: Any, source_side: str) -> list[str]:
    if source_side == "our_company":
        number = value(item, "internal_case_number") or value(item, "partner_case_number")
        received = value(item, "receipt_date_our_company")
        transferred = value(item, "transfer_date_our_company")
    else:
        number = value(item, "partner_case_number") or value(item, "partner_case_id") or value(item, "internal_case_number")
        received = value(item, "receipt_date_partner")
        transferred = value(item, "transfer_date_partner")
    description = value(item, "adverse_event") or value(item, "short_description") or "-"
    return [
        str(number or "-"),
        str(description),
        format_date(received),
        format_date(transferred),
    ]


def format_date(value_to_format: Any) -> str:
    if not value_to_format:
        return "-"
    if isinstance(value_to_format, str):
        try:
            value_to_format = date.fromisoformat(value_to_format[:10])
        except ValueError:
            return value_to_format
    if isinstance(value_to_format, date):
        return value_to_format.strftime("%d.%m.%Y")
    return str(value_to_format)


def value(item: Any, key: str) -> Any:
    if isinstance(item, dict):
        return item.get(key)
    return getattr(item, key, None)
